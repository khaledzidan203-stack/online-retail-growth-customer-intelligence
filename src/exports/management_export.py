"""Reproducible SQL -> pandas -> openpyxl exports; no SQL/Power BI writes.

Run: python src/exports/management_export.py
Revalidate saved files against current SQL: add --validate-only
Connection options: --server, --driver (Windows authentication only).
Amounts are source-currency units, not an asserted ISO currency.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone, date
from decimal import Decimal
import json
import math
from pathlib import Path
import re

import pandas as pd
import pyodbc
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import DataBarRule, ColorScaleRule
from openpyxl.utils import get_column_letter

import sql_queries as sql

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "outputs" / "management"
NAMES = ["Online_Retail_Management_Analytics.xlsx", "Online_Retail_Analytical_Detail.xlsx"]
REFERENCES = {
    "SalesRevenue": (19701685.507, .001), "UnitsSold": (11221960, 0),
    "Orders": (39519, 0), "SalesCustomers": (5852, 0),
    "CancellationValue": (719692.94, .001), "CancellationValueRate": (.035242, .000001),
    "AnonymousSalesRevenue": (2576013.46, .001), "OperationalAdjustmentRows": (3392, 0),
    "CanonicalTransactionRows": (1044848, 0),
}
EXEC = ["SalesRevenue", "UnitsSold", "Orders", "SalesCustomers", "AverageOrderValue",
        "RepeatCustomers", "OneTimeCustomers", "RepeatCustomerRate", "CancellationValue",
        "CancellationValueRate", "AnonymousSalesRevenue", "AnonymousSalesRevenuePct",
        "DuplicateFlaggedRevenue", "DuplicateRevenueImpactPct", "OperationalAdjustmentRows",
        "DQReviewRows"]
CAVEAT = ("Duplicate-flagged rows remain in official totals; anonymous activity can be valid. "
          "2009 and 2011 are partial periods. Distinct customers/orders are not additive across groups.")
DEFS = {
    "SalesRevenue": ("Approved merchandise sales", "SUM(LineAmount), IsSalesEligible=1"),
    "UnitsSold": ("Approved merchandise quantity", "SUM(Quantity), IsSalesEligible=1"),
    "Orders": ("Distinct eligible invoices", "COUNT(DISTINCT Invoice), IsSalesEligible=1"),
    "SalesCustomers": ("Known customers with eligible sales", "COUNT(DISTINCT CustomerKey), IsCustomerAnalyticsEligible=1"),
    "AverageOrderValue": ("Sales value per eligible order", "SalesRevenue / Orders"),
    "RepeatCustomers": ("Known customers with repeat purchase", "COUNT(CustomerRepeatBehavior), HasRepeatPurchase=1"),
    "OneTimeCustomers": ("Known customers without repeat purchase", "COUNT(CustomerRepeatBehavior), HasRepeatPurchase=0"),
    "RepeatCustomerRate": ("Share of known sales customers who repeat", "RepeatCustomers / (RepeatCustomers + OneTimeCustomers)"),
    "CancellationValue": ("Absolute approved customer cancellation value", "SUM(ABS(LineAmount)), IsCancellationEligible=1"),
    "CancellationValueRate": ("Cancellation share of gross sales and cancellation exposure", "CancellationValue / (SalesRevenue + CancellationValue)"),
    "AnonymousSalesRevenue": ("Eligible sales without known customer", "SUM(LineAmount), IsSalesEligible=1 AND IsKnownCustomer=0"),
    "AnonymousSalesRevenuePct": ("Anonymous share of eligible sales", "AnonymousSalesRevenue / SalesRevenue"),
    "DuplicateFlaggedRevenue": ("Sales exposure on duplicate-after-first rows", "SUM(LineAmount), IsSalesEligible=1 AND IsExactDuplicateAfterFirst=1"),
    "DuplicateRevenueImpactPct": ("Duplicate-after-first share of official sales", "DuplicateFlaggedRevenue / SalesRevenue"),
    "OperationalAdjustmentRows": ("Operational quantity-event rows", "COUNT(*), TransactionClass=OPERATIONAL_STOCK_ADJUSTMENT"),
    "DQReviewRows": ("Rows explicitly classified for review", "COUNT(*), TransactionClass=DQ_REVIEW"),
    "CanonicalTransactionRows": ("Preserved canonical lines", "COUNT(*) FROM FactTransaction"),
    "SalesEligibleRows": ("Approved merchandise sale lines", "COUNT(*), IsSalesEligible=1"),
    "UnknownCustomerRows": ("Lines with unknown customer identity", "COUNT(*), IsKnownCustomer=0"),
    "ZeroPriceRows": ("Zero-price exposure", "COUNT(*), IsZeroPrice=1"),
    "NegativePriceRows": ("Negative-price exposure", "COUNT(*), IsNegativePrice=1"),
    "ExactDuplicateAfterFirstRows": ("Duplicates after first occurrence", "COUNT(*), IsExactDuplicateAfterFirst=1"),
    "RowsInDuplicateGroups": ("All rows in exact duplicate groups", "COUNT(*), IsInExactDuplicateGroup=1"),
    "AccountingAdjustmentRows": ("Accounting adjustment lines", "COUNT(*), TransactionClass=ACCOUNTING_ADJUSTMENT"),
    "CancellationUnits": ("Absolute approved cancellation quantity", "SUM(ABS(Quantity)), IsCancellationEligible=1"),
    "CancellationInvoices": ("Distinct approved cancellation invoices", "COUNT(DISTINCT Invoice), IsCancellationEligible=1"),
    "CustomersWithCancellations": ("Known customers with approved cancellation", "COUNT(DISTINCT CustomerKey), IsCancellationEligible=1 AND IsKnownCustomer=1"),
    "RetentionRate": ("Observed cohort-period customer retention", "ActiveCustomers / CohortSize; use stored CohortRetention results"),
    "AverageFrequency": ("Mean published customer RFM order count", "AVG(CustomerRFM.Frequency) by Segment"),
    "AverageRecency": ("Mean days since last eligible purchase at reference date", "AVG(CustomerRFM.Recency) by Segment"),
    "AverageRevenuePerCustomer": ("Mean published RFM monetary value", "AVG(CustomerRFM.Monetary) by Segment"),
}
def query(c, statement):
    cur = c.cursor().execute(statement)
    return pd.DataFrame.from_records([tuple(r) for r in cur.fetchall()],
                                    columns=[x[0] for x in cur.description])

def close(a, b, tolerance=1e-7):
    return abs(float(a) - float(b)) <= tolerance

def metric_frame(values, keys):
    return pd.DataFrame([{"KPI": k, "Value": values[k],
                          "Definition": DEFS.get(k, (k, ""))[0]} for k in keys])

def collect(c):
    k = query(c, sql.KPI).iloc[0].to_dict()
    dq = query(c, sql.DQ).iloc[0].to_dict()
    k.update(dq)
    k["SalesCustomers"] = k["SalesKnownCustomers"]
    repeat = query(c, """SELECT SUM(CASE WHEN HasRepeatPurchase=1 THEN 1 ELSE 0 END) RepeatCustomers,
                        SUM(CASE WHEN HasRepeatPurchase=0 THEN 1 ELSE 0 END) OneTimeCustomers
                        FROM analytics.CustomerRepeatBehavior""").iloc[0].to_dict()
    assert all(k[n] == repeat[n] for n in repeat), "Repeat-model / KPI-view disagreement"
    k.update(repeat)
    # Ratios calculated from SQL aggregates, not rounded view ratios.
    for dest, num, den in [
        ("RepeatCustomerRate", "RepeatCustomers", "SalesCustomers"),
        ("AnonymousSalesRevenuePct", "AnonymousSalesRevenue", "SalesRevenue"),
        ("DuplicateRevenueImpactPct", "DuplicateFlaggedRevenue", "SalesRevenue")]:
        k[dest] = Decimal(str(k[num])) / Decimal(str(k[den])) if k[den] else None
    k["CancellationValueRate"] = k["CancellationValue"] / (k["SalesRevenue"] + k["CancellationValue"])
    independent = query(c, sql.INDEPENDENT).iloc[0].to_dict()
    independent["CancellationValueRate"] = independent["CancellationValue"] / (independent["SalesRevenue"] + independent["CancellationValue"])
    for n, (expected, tolerance) in REFERENCES.items():
        assert close(independent[n], expected, tolerance), f"STOP: reference discrepancy for {n}: {independent[n]} vs {expected}"
        assert close(k[n], independent[n]), f"STOP: SQL view/fact disagreement for {n}"
    frames = {n: query(c, q) for n, q in sql.QUERIES.items()}
    for n in ["MONTHLY_SALES", "MARKET_ANALYTICS", "PRODUCT_ANALYTICS"]:
        for metric in ["SalesRevenue", "UnitsSold", "CancellationValue"]:
            assert close(frames[n][metric].sum(), k[metric], .0001), f"Non-reconciling {n}/{metric}"
    period = query(c, "SELECT MIN(InvoiceDate) FirstDate,MAX(InvoiceDate) LastDate FROM analytics.FactTransaction").iloc[0]
    rel = query(c, sql.RELATIONSHIPS)
    assert len(rel) == 7, "Unexpected relationship count"
    assert len(frames["CUSTOMER_360"]) == k["SalesCustomers"]
    segments = query(c, sql.RFM_SEGMENTS)
    customers = query(c, sql.CUSTOMERS)
    classes = frames["TRANSACTION_CLASS_DETAIL"]
    assert int(classes["Rows"].sum()) == int(k["CanonicalTransactionRows"])
    adjustments = classes[classes.TransactionClass.isin(["OPERATIONAL_STOCK_ADJUSTMENT", "ACCOUNTING_ADJUSTMENT"])].copy()
    operational = adjustments[adjustments.TransactionClass == "OPERATIONAL_STOCK_ADJUSTMENT"]
    assert all(operational["AbsoluteValue"] == 0), "Operational adjustments carry unexpected money"
    dictionary = pd.DataFrame([{"KPI": n, "Definition": d, "CalculationLogic": logic,
                               "SourceTable": ("analytics.CustomerRepeatBehavior" if n in ["RepeatCustomers","OneTimeCustomers","RepeatCustomerRate"]
                                               else "analytics.CohortRetention" if n == "RetentionRate"
                                               else "analytics.CustomerRFM" if n.startswith("AverageR") or n == "AverageFrequency"
                                               else "analytics.FactTransaction"),
                               "BusinessMeaning": d, "ImportantCaveat": CAVEAT}
                              for n, (d, logic) in DEFS.items()])
    # Each sheet is a list of labelled tables; cancellations include four complementary grains.
    management = {
        "01_EXECUTIVE_SUMMARY": [("Full-period governed KPIs", metric_frame(k, EXEC))],
        "02_MONTHLY_SALES": [("Calendar month", frames["MONTHLY_SALES"])],
        "03_MARKET_PERFORMANCE": [("Recorded transaction country", frames["MARKET_ANALYTICS"])],
        "04_PRODUCT_PERFORMANCE": [("Merchandise products ranked by revenue", frames["PRODUCT_ANALYTICS"].head(100))],
        "05_CUSTOMER_PERFORMANCE": [("Known sales customers", customers)],
        "06_RFM_SEGMENTS": [("Published snapshot segment", segments)],
        "07_COHORT_RETENTION": [("Published cohort-period cells; flags retained", frames["COHORT_RETENTION"])],
        "08_CANCELLATIONS": [
            ("Full-period cancellation KPIs", metric_frame(k, ["CancellationValue","CancellationUnits","CancellationInvoices","CustomersWithCancellations","CancellationValueRate"])),
            ("Monthly cancellations", frames["MONTHLY_CANCELLATIONS"]),
            ("Top 20 products by cancellation value", frames["CANCELLATION_PRODUCTS"].head(20)),
            ("Top 20 markets by cancellation value", frames["CANCELLATION_MARKETS"].head(20))],
        "09_ADJUSTMENTS": [("Separate operational and accounting classes", adjustments)],
        "10_DATA_QUALITY": [("Governance and exposure, not automatic invalidity",
                            metric_frame(k, list(dq) + ["AnonymousSalesRevenuePct","DuplicateRevenueImpactPct"]))],
        "11_TRANSACTION_CLASSES": [("Canonical class profile", classes)],
        "12_KPI_DICTIONARY": [("Governed definitions", dictionary)],
        "13_MODEL_RELATIONSHIPS": [("Actual enabled SQL foreign keys; not extra report relationships", rel)],
    }
    detail = {name: [("Existing SQL analytics / governed summary", frame)] for name, frame in frames.items()}
    generated = datetime.now(timezone.utc).isoformat(timespec="seconds")
    for sheets, purpose in [(management, "Management review; top 100 products, all known sales customers"),
                             (detail, "Analyst drill-down; all published analytical rows, never the raw fact")]:
        items = [
            ("Project", "Online Retail Growth & Customer Intelligence"),
            ("Dataset", "UCI Online Retail II"),
            ("Analysis period", f"{period.FirstDate:%Y-%m-%d} to {period.LastDate:%Y-%m-%d}"),
            ("Analytical source", "SQL Server analytics model"),
            ("Generated by", "Programmatic analytical export pipeline"),
            ("Generated timestamp UTC", generated), ("Purpose", purpose),
            ("Scope / grain", "Full period; sheet labels identify customer, product, country, month and cohort-period grains."),
            ("Sales", DEFS["SalesRevenue"][1]), ("Cancellation rate", DEFS["CancellationValueRate"][1]),
            ("Currency", "SC denotes source-currency units; no ISO currency is asserted. Display rounding does not change stored values."),
            ("Governance", CAVEAT),
            ("RFM/cohorts", "Published snapshot RFM is not dynamically re-segmented; preserve left-boundary, partial-observation and cohort indexing flags."),
            ("Relationships", "Actual SQL foreign keys shown. SQL customer tables are unique per customer; Power BI uses seven governed single-direction M:1 relationships. CohortRetention and CohortRevenue are disconnected."),
            ("Aggregation", "Do not sum rates, customer counts, cumulative shares, ranks or repeated cohort sizes. Cohort retention across cells is weighted by cohort-size exposure."),
            ("Cancellation risk", "Keep outliers. Interpret low-volume rates cautiously; country view includes the existing minimum-base flag."),
            ("Supported metrics", "No profit, margin, COGS, inventory value, ROI or marketing-spend estimates."),
            ("Classifications", ", ".join(classes.TransactionClass)),
        ]
        items.extend((name, " / ".join(label for label, _ in blocks)) for name, blocks in sheets.items())
        sheets["00_README"] = [("Workbook guide", pd.DataFrame(items, columns=["Item", "Detail"]))]
    return management, detail, k, generated

def number_format(name):
    if name.endswith("Rank") or name == "Year": return "0"
    if any(x in name for x in ["Rate", "Share", "Pct", "Impact"]): return "0.00%"
    if any(x in name for x in ["Date", "Month", "Cohort", "PeriodStart"]) and name not in ["MonthNumber","CohortIndex","CohortSize","YearMonth"]: return "yyyy-mm-dd"
    if any(x in name for x in ["Revenue","Monetary","Value","Price"]): return '"SC " #,##0.00;[Red]("SC " #,##0.00);"-"'
    if name.startswith("Average") or "PerOriginalCustomer" in name: return "#,##0.00"
    return "#,##0"

def scalar(value):
    if pd.isna(value): return None
    if isinstance(value, pd.Timestamp): return value.to_pydatetime()
    if isinstance(value, Decimal): return float(value)
    return value.item() if hasattr(value, "item") else value

def write_workbook(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "Analytical Export Pipeline"
    wb.properties.title = "Online Retail Growth & Customer Intelligence"
    table_no = 0
    for name in ["00_README"] + [s for s in sheets if s != "00_README"]:
        ws = wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "B5"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        row = 1
        for label, df in sheets[name]:
            assert len(df) + row + 4 <= 1048576
            ws.cell(row, 1, name.replace("_", " "))
            ws.cell(row, 1).font = Font(name="Calibri", size=17, bold=True, color="FFFFFF")
            ws.cell(row, 1).fill = PatternFill("solid", fgColor="142338")
            ws.cell(row + 1, 1, label).font = Font(size=11, italic=True, color="596579")
            header = row + 3
            for col, column in enumerate(df.columns, 1):
                cell = ws.cell(header, col, column)
                cell.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
                cell.fill = PatternFill("solid", fgColor="142338")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[header].height = 32
            for offset, record in enumerate(df.itertuples(index=False, name=None), header + 1):
                for col, (column, value) in enumerate(zip(df.columns, record), 1):
                    cell = ws.cell(offset, col, scalar(value))
                    if isinstance(cell.value, str): cell.data_type = "s"  # neutralize Excel formula injection
                    cell.font = Font(name="Calibri", size=11)
                    metric = record[0] if column == "Value" and "KPI" in df else column
                    cell.number_format = number_format(str(metric))
                    cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str))
                if name in ["00_README","12_KPI_DICTIONARY"]: ws.row_dimensions[offset].height = 45
            end = header + len(df)
            table_no += 1
            if len(df):
                tab = Table(displayName=f"ExportTable{table_no}", ref=f"A{header}:{get_column_letter(len(df.columns))}{end}")
                tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
                ws.add_table(tab)
                for i, column in enumerate(df.columns, 1):
                    cells = f"{get_column_letter(i)}{header+1}:{get_column_letter(i)}{end}"
                    if column in ["SalesRevenue","CancellationValue","Monetary"]:
                        ws.conditional_formatting.add(cells, DataBarRule(start_type="num",start_value=0,end_type="max",
                            color="D64554" if column == "CancellationValue" else "20A99B",showValue=True))
                    if column == "RetentionRate":
                        ws.conditional_formatting.add(cells, ColorScaleRule(start_type="num",start_value=0,start_color="FFFFFF",
                                                                          end_type="num",end_value=1,end_color="20A99B"))
            for i, column in enumerate(df.columns, 1):
                lengths = [len(str(x)) for x in df[column].head(100) if pd.notna(x)]
                width = min(44, max(16, len(column)+2, min(max(lengths, default=0)+2, 35)))
                if column in ["Detail","ImportantCaveat","CalculationLogic"]: width = 70
                if column in ["Definition","BusinessMeaning","RepresentativeDescription"]: width = 44
                letter = get_column_letter(i)
                ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, width)
            row = end + 4
        ws.print_title_rows = "1:4"
    wb.save(path)
    wb.close()

def validate_workbook(path, sheets):
    """Compare every exported cell and table range to fresh SQL-derived frames."""
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["00_README"] + [s for s in sheets if s != "00_README"]
    count = 0
    for name, blocks in sheets.items():
        ws = wb[name]
        assert ws.freeze_panes == "B5"
        row = 1
        for _, df in blocks:
            header = row + 3
            assert [ws.cell(header,i).value for i in range(1,len(df.columns)+1)] == list(df.columns)
            if len(df):
                ref = f"A{header}:{get_column_letter(len(df.columns))}{header+len(df)}"
                assert any(t.ref == ref for t in ws.tables.values()), f"Table missing: {name}"
            for r, record in enumerate(df.itertuples(index=False,name=None), header+1):
                for col, expected in enumerate(record,1):
                    expected = scalar(expected)
                    actual = ws.cell(r,col).value
                    assert ws.cell(r,col).data_type != "f", f"Unexpected formula: {name}"
                    # Timestamp intentionally changes on a new validation invocation.
                    if name == "00_README" and record[0] == "Generated timestamp UTC": continue
                    if isinstance(expected,(int,float)) and not isinstance(expected,bool):
                        assert actual is not None and close(actual,expected), f"Value mismatch {name}!{r},{col}"
                    elif isinstance(expected,date) and not isinstance(expected,datetime):
                        assert isinstance(actual,datetime) and actual.date() == expected, f"Date mismatch {name}"
                    elif isinstance(expected,datetime):
                        assert actual == expected, f"Date mismatch {name}"
                    else: assert actual == expected or actual is None and expected == "", f"Cell mismatch {name}!{r},{col}"
                    count += 1
            row = header + len(df) + 4
    wb.close()
    return count

def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--server", default="localhost")
    p.add_argument("--driver", default="ODBC Driver 17 for SQL Server")
    p.add_argument("--validate-only", action="store_true")
    args = p.parse_args()
    # Read-only statements only; rollback closes the consistent SERIALIZABLE read transaction.
    c = pyodbc.connect(f"DRIVER={{{args.driver}}};SERVER={args.server};DATABASE=OnlineRetailAnalytics;"
                       "Trusted_Connection=yes;Encrypt=yes;TrustServerCertificate=yes", autocommit=False)
    try:
        c.cursor().execute("SET TRANSACTION ISOLATION LEVEL SERIALIZABLE")
        management, detail, k, generated = collect(c)
    finally:
        c.rollback()
        c.close()
    OUT.mkdir(parents=True, exist_ok=True)
    results = {}
    for filename, sheets in zip(NAMES,[management,detail]):
        path = OUT / filename
        if not args.validate_only: write_workbook(path,sheets)
        cells = validate_workbook(path,sheets)
        results[filename] = {"sheets":len(sheets),"validated_cells":cells,
                             "sheet_rows":{n:sum(len(df) for _,df in blocks) for n,blocks in sheets.items()}}
    result = {"status":"PASS","generated_utc":generated,"sql_validation":"PASS","reference_reconciliation":"PASS",
              "references":{n:{"actual":str(k[n]),"reference":str(v[0]),"tolerance":v[1]} for n,v in REFERENCES.items()},
              "workbooks":results,"warnings":["Source-currency code is unspecified; amounts use SC formatting.",
                "2009 and 2011 are partial periods; cohort observation-boundary flags are retained.",
                "Distinct counts and rates are nonadditive; management product/cancellation lists are explicitly Top N."]}
    (OUT / "management_export_validation.json").write_text(json.dumps(result,indent=2),encoding="utf-8")
    print(json.dumps(result,indent=2))

if __name__ == "__main__":
    main()
